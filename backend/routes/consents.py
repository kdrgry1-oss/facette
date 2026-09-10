"""
PAZARLAMA İZİNLERİ (E-posta & SMS) — kim izin verdi / kim reddetti, tek listede.

Kaynaklar (birleştirilir; alıcı+kanal başına EN SON kayıt geçerli):
  • db.iys_consents          — sipariş/üyelik/OTP akışlarından İYS izin kayıtları
                               (channels: MESAJ=SMS, EPOSTA=e-posta; status: ONAY / RET)
  • db.newsletter_subscribers — site altı bülten aboneliği (e-posta; consent=True, active)
  • db.email_suppressions    — bounce/şikâyet kara listesi (e-posta izni fiilen geçersiz)

GET  /admin/consents?channel=email|sms&status=onay|ret&q=&page=&limit=
GET  /admin/consents/summary
GET  /admin/consents/export.xlsx
"""
import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from .deps import db, require_admin

router = APIRouter(prefix="/admin/consents", tags=["admin-consents"])

_SRC_TR = {"HS_WEB": "Web sitesi", "HS_FIZIKSEL_ORTAM": "Fiziksel ortam", "HS_MESAJ": "SMS ile", "HS_EPOSTA": "E-posta ile",
           "footer": "Bülten (site altı)", "checkout": "Ödeme sayfası", "register": "Üyelik", "otp": "SMS doğrulama",
           "popup": "Açılır pencere", "account": "Hesabım"}


async def _collect() -> dict:
    """{(kanal, alıcı): kayıt} — en güncel kayıt kazanır."""
    latest: dict = {}

    def _put(channel, recipient, rec):
        if not recipient:
            return
        key = (channel, recipient)
        cur = latest.get(key)
        if not cur or (rec.get("at") or "") >= (cur.get("at") or ""):
            latest[key] = rec

    async for r in db.iys_consents.find({}, {"_id": 0}).sort("created_at", 1):
        at = r.get("consent_date") or r.get("created_at") or ""
        st = "onay" if str(r.get("status") or "ONAY").upper() == "ONAY" else "ret"
        base = {"at": at, "status": st, "source": _SRC_TR.get(str(r.get("source") or ""), str(r.get("source") or "")),
                "origin": "İYS kaydı", "email": (r.get("email") or "").strip().lower(),
                "phone": r.get("phone") or "", "order_id": r.get("order_id") or "", "user_id": r.get("user_id"),
                "reported": bool(r.get("reported"))}
        for ch in (r.get("channels") or []):
            if ch == "EPOSTA" and base["email"]:
                _put("email", base["email"], dict(base, recipient=base["email"]))
            elif ch == "MESAJ" and base["phone"]:
                _put("sms", base["phone"], dict(base, recipient=base["phone"]))

    async for s in db.newsletter_subscribers.find({}, {"_id": 0}):
        em = (s.get("email") or "").strip().lower()
        if not em:
            continue
        at = s.get("unsubscribed_at") or s.get("consent_at") or s.get("created_at") or ""
        st = "onay" if (s.get("consent") and s.get("active", True) is not False) else "ret"
        _put("email", em, {"at": at, "status": st, "source": _SRC_TR.get(str(s.get("source") or "footer"), str(s.get("source") or "")),
                          "origin": "Bülten aboneliği", "email": em, "phone": "", "recipient": em,
                          "order_id": "", "user_id": None, "reported": True})

    suppressed = set()
    async for x in db.email_suppressions.find({}, {"_id": 0, "email": 1, "reason": 1}):
        em = (x.get("email") or "").strip().lower()
        if em:
            suppressed.add(em)
    for (ch, rcp), rec in latest.items():
        if ch == "email" and rcp in suppressed:
            rec["suppressed"] = True
    return latest


async def _names_for(recs: list) -> None:
    """Üye adını e-posta/telefon eşleşmesiyle doldur (tek sorgu)."""
    emails = {r["recipient"] for r in recs if r.get("channel") == "email"}
    phones = {r["recipient"] for r in recs if r.get("channel") == "sms"}
    by_email, by_phone = {}, {}
    if emails:
        async for u in db.users.find({"email": {"$in": list(emails)}}, {"_id": 0, "email": 1, "first_name": 1, "last_name": 1, "phone": 1}):
            by_email[(u.get("email") or "").lower()] = u
    if phones:
        import re as _re
        tails = [p[-10:] for p in phones if len(p) >= 10]
        if tails:
            async for u in db.users.find({"phone": {"$regex": "(" + "|".join(_re.escape(t) for t in tails) + ")$"}},
                                         {"_id": 0, "email": 1, "first_name": 1, "last_name": 1, "phone": 1}):
                d = _re.sub(r"\D", "", str(u.get("phone") or ""))
                if len(d) >= 10:
                    by_phone[d[-10:]] = u
    for r in recs:
        u = by_email.get(r["recipient"]) if r.get("channel") == "email" else by_phone.get(r["recipient"][-10:])
        if u:
            r["name"] = f"{u.get('first_name', '')} {u.get('last_name', '')}".strip()
            r["member_email"] = u.get("email")


def _rows(latest: dict, channel: Optional[str], status: Optional[str], q: str) -> list:
    out = []
    ql = (q or "").strip().lower()
    for (ch, rcp), rec in latest.items():
        if channel and ch != channel:
            continue
        st = "ret" if rec.get("suppressed") else rec.get("status")
        if status and st != status:
            continue
        if ql and ql not in rcp and ql not in (rec.get("email") or "") and ql not in (rec.get("phone") or ""):
            continue
        out.append({"channel": ch, "recipient": rcp, "status": st, "suppressed": bool(rec.get("suppressed")),
                    "at": rec.get("at"), "source": rec.get("source"), "origin": rec.get("origin"),
                    "email": rec.get("email"), "phone": rec.get("phone"), "reported": rec.get("reported")})
    out.sort(key=lambda r: r.get("at") or "", reverse=True)
    return out


@router.get("/summary")
async def consents_summary(current_user: dict = Depends(require_admin)):
    latest = await _collect()
    def _n(ch, st):
        return sum(1 for (c, _), r in latest.items() if c == ch and (("ret" if r.get("suppressed") else r.get("status")) == st))
    return {"email": {"onay": _n("email", "onay"), "ret": _n("email", "ret")},
            "sms": {"onay": _n("sms", "onay"), "ret": _n("sms", "ret")},
            "generated_at": datetime.now(timezone.utc).isoformat()}


@router.get("")
async def list_consents(channel: Optional[str] = Query(None), status: Optional[str] = Query(None),
                        q: str = "", page: int = Query(1, ge=1), limit: int = Query(50, ge=1, le=500),
                        current_user: dict = Depends(require_admin)):
    latest = await _collect()
    rows = _rows(latest, channel, status, q)
    total = len(rows)
    page_rows = rows[(page - 1) * limit: page * limit]
    await _names_for(page_rows)
    return {"items": page_rows, "total": total, "page": page, "pages": max(1, (total + limit - 1) // limit)}


@router.get("/export.xlsx")
async def export_consents(channel: Optional[str] = Query(None), status: Optional[str] = Query(None), q: str = "",
                          current_user: dict = Depends(require_admin)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    latest = await _collect()
    rows = _rows(latest, channel, status, q)
    await _names_for(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "İzinler"
    ws.append(["Kanal", "Alıcı", "Ad Soyad", "Durum", "Kaynak", "Kayıt türü", "Tarih", "E-posta", "Telefon", "İYS'ye bildirildi", "Kara liste"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="111111")
    for r in rows:
        ws.append(["E-posta" if r["channel"] == "email" else "SMS", r["recipient"], r.get("name") or "",
                   "İZİN VERDİ" if r["status"] == "onay" else "RET", r.get("source") or "", r.get("origin") or "",
                   str(r.get("at") or "")[:16].replace("T", " "), r.get("email") or "", r.get("phone") or "",
                   "Evet" if r.get("reported") else "Hayır", "Evet" if r.get("suppressed") else ""])
    for i, w in enumerate([10, 30, 24, 12, 18, 18, 17, 30, 16, 12, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"pazarlama-izinleri-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}", "Cache-Control": "no-store"})
